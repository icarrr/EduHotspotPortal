from flask import Blueprint, render_template, request, flash, redirect, url_for, make_response
from flask_login import login_required, current_user
from app.decorators import operator_required, admin_required
from app.models import HotspotUser, Operator, AuditLog, db, RoleProfile
from app.utils.mikrotik import get_mikrotik_client
import openpyxl
import io
import secrets
import string

bp = Blueprint('users', __name__, url_prefix='/users')


@bp.route('/')
@login_required
@operator_required
def index():
    search = request.args.get('search', '').strip().lower()
    mikrotik = get_mikrotik_client()
    hotspot_users = mikrotik.get_hotspot_users()

    if search:
        hotspot_users = [u for u in hotspot_users
                        if search in str(u.get('name', '')).lower() or
                        search in str(u.get('comment', '')).lower()]
        if not hotspot_users:
            flash(f'No users found for "{search}"', 'info')

    return render_template('users/index.html', users=hotspot_users, search=search)


@bp.route('/add', methods=['GET', 'POST'])
@login_required
@operator_required
def add():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password') or generate_password()
        role = request.form.get('role', 'siswa')
        profile = role

        mikrotik = get_mikrotik_client()
        success, message = mikrotik.add_hotspot_user(username, password, profile, role)

        if success:
            flash(f'User {username} added successfully. Password: {password}', 'success')
            log = AuditLog(operator_id=current_user.id, action='create',
                          target_user=username, details=f'Role: {role}')
            db.session.add(log)
            db.session.commit()
            return redirect(url_for('users.index'))
        else:
            flash(f'Failed to add user: {message}', 'danger')

    return render_template('users/add.html')


@bp.route('/edit/<username>', methods=['GET', 'POST'])
@login_required
@operator_required
def edit(username):
    mikrotik = get_mikrotik_client()
    user = None
    users = mikrotik.get_hotspot_users()
    for u in users:
        if str(u.get('name')) == str(username):
            user = u
            break

    if not user:
        flash('User not found', 'danger')
        return redirect(url_for('users.index'))

    if request.method == 'POST':
        profile = request.form.get('profile')
        disabled = 'no' if request.form.get('enabled') else 'yes'
        mikrotik.edit_hotspot_user(username, profile=profile, disabled=disabled)
        flash(f'User {username} updated', 'success')
        return redirect(url_for('users.index'))

    return render_template('users/edit.html', user=user)


@bp.route('/delete/<username>')
@login_required
@admin_required
def delete(username):
    mikrotik = get_mikrotik_client()
    success, message = mikrotik.delete_hotspot_user(username)
    if success:
        flash(f'User {username} deleted', 'success')
    else:
        flash(f'Failed to delete user: {message}', 'danger')
    return redirect(url_for('users.index'))


@bp.route('/bulk-delete', methods=['POST'])
@login_required
@admin_required
def bulk_delete():
    usernames = request.form.getlist('usernames')
    if not usernames:
        flash('No users selected', 'warning')
        return redirect(url_for('users.index'))

    mikrotik = get_mikrotik_client()
    deleted = 0
    failed = 0

    for username in usernames:
        success, message = mikrotik.delete_hotspot_user(username)
        if success:
            deleted += 1
            log = AuditLog(operator_id=current_user.id, action='delete',
                          target_user=username, details='Bulk delete')
            db.session.add(log)
        else:
            failed += 1

    db.session.commit()

    if deleted > 0:
        flash(f'{deleted} user(s) deleted successfully', 'success')
    if failed > 0:
        flash(f'{failed} user(s) failed to delete', 'danger')

    return redirect(url_for('users.index'))


@bp.route('/reset-password/<username>')
@login_required
@operator_required
def reset_password(username):
    new_password = secrets.token_urlsafe(8)
    mikrotik = get_mikrotik_client()
    success, message = mikrotik.reset_password(username, new_password)
    if success:
        flash(f'Password for {username} reset to: {new_password}', 'success')
    else:
        flash(f'Failed to reset password: {message}', 'danger')
    return redirect(url_for('users.index'))


@bp.route('/reset-counters/<username>')
@login_required
@operator_required
def reset_counters(username):
    mikrotik = get_mikrotik_client()
    success, message = mikrotik.reset_user_counters(username)
    if success:
        flash(f'Counters reset for {username}', 'success')
        log = AuditLog(operator_id=current_user.id, action='update',
                      target_user=username, details='Counters reset')
        db.session.add(log)
        db.session.commit()
    else:
        flash(f'Failed to reset counters: {message}', 'danger')
    return redirect(url_for('users.index'))


@bp.route('/toggle/<username>')
@login_required
@operator_required
def toggle(username):
    mikrotik = get_mikrotik_client()
    users = mikrotik.get_hotspot_users()
    for u in users:
        if str(u.get('name')) == str(username):
            if u.get('disabled') == True:
                mikrotik.enable_user(username)
                flash(f'User {username} enabled', 'success')
            else:
                mikrotik.disable_user(username)
                flash(f'User {username} disabled', 'success')
            break
    return redirect(url_for('users.index'))


@bp.route('/online')
@login_required
@operator_required
def online_sessions():
    mikrotik = get_mikrotik_client()
    active_sessions = mikrotik.get_active_sessions()
    hosts = mikrotik.get_hosts()

    logged_in_macs = {s.get('mac-address', '') for s in active_sessions}
    merged_sessions = []
    for h in hosts:
        mac = h.get('mac-address', '')
        for s in active_sessions:
            if s.get('mac-address', '') == mac:
                h['user'] = s.get('user')
                h['address'] = s.get('address')
                h['uptime'] = s.get('uptime')
                h['idle-time'] = s.get('idle-time')
                h['.id'] = s['.id']
                h['logged_in'] = True
                merged_sessions.append(h)
                break
        else:
            h['user'] = None
            h['logged_in'] = False
            merged_sessions.append(h)

    return render_template('users/online.html', sessions=merged_sessions)


@bp.route('/disconnect/<session_id>')
@login_required
@operator_required
def disconnect(session_id):
    mikrotik = get_mikrotik_client()
    success, message = mikrotik.disconnect_session(session_id)
    if success:
        flash('Session disconnected', 'success')
    else:
        flash(f'Failed to disconnect: {message}', 'danger')
    return redirect(url_for('users.online_sessions'))


@bp.route('/import', methods=['GET', 'POST'])
@login_required
@operator_required
def import_users():
    if request.method == 'POST':
        if 'xlsx_file' not in request.files:
            flash('No file uploaded', 'danger')
            return redirect(url_for('users.import_users'))

        file = request.files['xlsx_file']
        if file.filename == '':
            flash('No file selected', 'danger')
            return redirect(url_for('users.import_users'))

        if not file.filename.endswith('.xlsx'):
            flash('Please upload an XLSX file', 'danger')
            return redirect(url_for('users.import_users'))

        try:
            file_data = io.BytesIO(file.read())
            wb = openpyxl.load_workbook(file_data)
            ws = wb.active

            mikrotik = get_mikrotik_client()
            existing_users = mikrotik.get_hotspot_users()
            existing_usernames = {u.get('name') for u in existing_users}

            created = 0
            skipped = 0
            errors = []
            created_users = []

            # Assume first row is header: username, nama, role
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            for row in rows:
                if not row or not row[0]:
                    continue
                username = str(row[0]).strip() if row[0] else ''
                name = str(row[1]).strip() if row[1] else ''
                role = str(row[2]).strip().lower() if row[2] else 'siswa'

                # Validate role
                if role not in ['admin', 'guru', 'siswa', 'trial']:
                    role = 'siswa'

                if not username or username == 'None':
                    errors.append('Missing username in row')
                    continue

                if username in existing_usernames:
                    skipped += 1
                    continue

                password = generate_password()
                profile = role

                success, message = mikrotik.add_hotspot_user(username, password, profile, role)

                if success:
                    created += 1
                    existing_usernames.add(username)
                    created_users.append({'username': username, 'password': password, 'name': name})
                    log = AuditLog(operator_id=current_user.id, action='create',
                                  target_user=username, details=f'Imported via XLSX, Role: {role}, Password: {password}')
                    db.session.add(log)
                else:
                    errors.append(f'{username}: {message}')

            db.session.commit()

            if created_users:
                return render_template('users/import_results.html', created_users=created_users,
                                      skipped=skipped, errors=errors)
            else:
                if skipped:
                    flash(f'Skipped {skipped} duplicate users', 'warning')
                if errors:
                    for err in errors[:5]:
                        flash(err, 'danger')
                return redirect(url_for('users.index'))

        except Exception as e:
            flash(f'Error processing file: {str(e)}', 'danger')
            return redirect(url_for('users.import_users'))

    return render_template('users/import.html')


@bp.route('/operators', methods=['GET', 'POST'])
@login_required
@admin_required
def operators():
    if request.method == 'POST':
        username = request.form.get('username')
        name = request.form.get('name')
        password = request.form.get('password')
        role = request.form.get('role', 'operator')

        if not username or not password:
            flash('Username and password required', 'danger')
            return redirect(url_for('users.operators'))

        if role not in ['admin_it', 'operator']:
            flash('Invalid role', 'danger')
            return redirect(url_for('users.operators'))

        existing = Operator.query.filter_by(username=username).first()
        if existing:
            flash('Username already exists', 'danger')
            return redirect(url_for('users.operators'))

        op = Operator(username=username, name=name, role=role, active=True)
        op.set_password(password)
        db.session.add(op)
        db.session.commit()

        log = AuditLog(operator_id=current_user.id, action='create',
                      target_user=username, details=f'New operator, Role: {role}')
        db.session.add(log)
        db.session.commit()

        flash(f'Operator {username} added successfully', 'success')
        return redirect(url_for('users.operators'))

    operators = Operator.query.all()
    return render_template('users/operators.html', operators=operators)


@bp.route('/audit-logs')
@login_required
@admin_required
def audit_logs():
    page = request.args.get('page', 1, type=int)
    logs = AuditLog.query.order_by(AuditLog.timestamp.desc()).paginate(page=page, per_page=50)
    return render_template('users/audit_logs.html', logs=logs)


@bp.route('/download-template')
@login_required
@operator_required
def download_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Users Template'
    ws.append(['username', 'nama', 'role'])
    ws.append(['12345', 'Budi Santoso', 'siswa'])
    ws.append(['67890', 'Siti Nurhaliza', 'guru'])
    ws.append(['11111', 'Ahmad Staff', 'trial'])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = make_response(output.read())
    response.headers['Content-Disposition'] = 'attachment; filename=users_template.xlsx'
    response.headers['Content-type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


@bp.route('/download-passwords')
@login_required
@operator_required
def download_passwords():
    import json
    data = request.args.get('data', '[]')
    try:
        users = json.loads(data)
    except:
        users = []

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'User Passwords'
    ws.append(['username', 'nama', 'password'])

    for user in users:
        ws.append([user.get('username', ''), user.get('name', ''), user.get('password', '')])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = make_response(output.read())
    response.headers['Content-Disposition'] = 'attachment; filename=user_passwords.xlsx'
    response.headers['Content-type'] = 'application/vnd.openxmlformats-officedocument.spspreadsheetml.sheet'
    return response


@bp.route('/export')
@login_required
@operator_required
def export_users():
    mikrotik = get_mikrotik_client()
    hotspot_users = mikrotik.get_hotspot_users()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Hotspot Users'
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 15

    from openpyxl.styles import numbers
    ws.column_dimensions['A'].number_format = numbers.FORMAT_TEXT
    ws.column_dimensions['B'].number_format = numbers.FORMAT_TEXT

    ws.append(['username', 'password', 'role', 'profile'])

    for user in hotspot_users:
        username = str(user.get('name', ''))
        password = str(user.get('password', ''))
        comment = user.get('comment', '')
        role = comment.replace('role:', '') if comment.startswith('role:') else ''
        profile = user.get('profile', '')

        row_num = ws.max_row + 1
        ws.cell(row=row_num, column=1, value=username).number_format = numbers.FORMAT_TEXT
        ws.cell(row=row_num, column=2, value=password).number_format = numbers.FORMAT_TEXT
        ws.cell(row=row_num, column=3, value=role)
        ws.cell(row=row_num, column=4, value=profile)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = make_response(output.read())
    response.headers['Content-Disposition'] = 'attachment; filename=hotspot_users.xlsx'
    response.headers['Content-type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


@bp.route('/bind-mac/<username>', methods=['GET', 'POST'])
@login_required
@operator_required
def bind_mac(username):
    """Bind MAC address to user for device binding"""
    from app.models import HotspotUser

    # Get or create local user record
    local_user = HotspotUser.query.filter_by(username=username).first()
    if not local_user:
        local_user = HotspotUser(username=username)
        db.session.add(local_user)

    if request.method == 'POST':
        mac_address = request.form.get('mac_address', '').strip().upper()

        # Validate MAC address format
        import re
        mac_pattern = re.compile(r'^([0-9A-F]{2}[:-]){5}([0-9A-F]{2})$')
        if not mac_pattern.match(mac_address):
            flash('Invalid MAC address format. Use AA:BB:CC:DD:EE:FF', 'danger')
            return redirect(url_for('users.bind_mac', username=username))

        mikrotik = get_mikrotik_client()
        success, message = mikrotik.bind_mac_to_user(username, mac_address)

        if success:
            local_user.mac_address = mac_address
            db.session.commit()
            flash(f'MAC address {mac_address} bound to {username}', 'success')
            log = AuditLog(operator_id=current_user.id, action='update',
                          target_user=username, details=f'MAC bound: {mac_address}')
            db.session.add(log)
            db.session.commit()
            return redirect(url_for('users.index'))
        else:
            flash(f'Failed to bind MAC: {message}', 'danger')

    return render_template('users/bind_mac.html', username=username,
                         current_mac=local_user.mac_address)


@bp.route('/unbind-mac/<username>')
@login_required
@operator_required
def unbind_mac(username):
    """Remove MAC address binding from user"""
    from app.models import HotspotUser

    local_user = HotspotUser.query.filter_by(username=username).first()
    if local_user:
        local_user.mac_address = None
        db.session.commit()

    mikrotik = get_mikrotik_client()
    success, message = mikrotik.unbind_mac_from_user(username)

    if success:
        flash(f'MAC address unbound from {username}', 'success')
        log = AuditLog(operator_id=current_user.id, action='update',
                      target_user=username, details='MAC unbound')
        db.session.add(log)
        db.session.commit()
    else:
        flash(f'Failed to unbind MAC: {message}', 'danger')

    return redirect(url_for('users.index'))


@bp.route('/set-expiration/<username>', methods=['GET', 'POST'])
@login_required
@operator_required
def set_expiration(username):
    """Set expiration date for user"""
    from app.models import HotspotUser
    from datetime import datetime

    local_user = HotspotUser.query.filter_by(username=username).first()
    if not local_user:
        local_user = HotspotUser(username=username)
        db.session.add(local_user)

    if request.method == 'POST':
        expire_date_str = request.form.get('expire_date', '')
        if expire_date_str:
            try:
                expire_date = datetime.strptime(expire_date_str, '%Y-%m-%d')
                local_user.expires_at = expire_date
                db.session.commit()
                flash(f'Expiration set for {username}', 'success')
                log = AuditLog(operator_id=current_user.id, action='update',
                              target_user=username, details=f'Expires: {expire_date_str}')
                db.session.add(log)
                db.session.commit()
            except ValueError:
                flash('Invalid date format', 'danger')
        else:
            local_user.expires_at = None
            db.session.commit()
            flash(f'Expiration removed for {username}', 'success')

        return redirect(url_for('users.index'))

    return render_template('users/set_expiration.html', username=username,
                         current_expires=local_user.expires_at)


def generate_password(length=6):
    return ''.join(secrets.choice(string.ascii_letters + string.digits) for _ in range(length))
